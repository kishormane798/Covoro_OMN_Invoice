from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

_UTILS = Path(__file__).resolve().parent
if str(_UTILS) not in sys.path:
    sys.path.insert(0, str(_UTILS))

from write_copy_validation_matrix import (  # noqa: E402
    COPY_INVOICE_PRECONDITIONS,
    copy_invoice_steps,
    copy_kept_rows,
    rewrite_copy_invoice_narrative,
    write_combined_workbook,
)


def _matrix_xlsx(path: Path, rows: list[tuple[str, str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Testcases"
    ws.append(["Test Case ID", "Filed name", "Testcase Title"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    wb.close()


class RewriteCopyInvoiceNarrativeTests(unittest.TestCase):
    def test_rewrites_excel_upload_phrases_to_copy_form(self) -> None:
        rewritten = rewrite_copy_invoice_narrative(
            "Excel upload · Covoro | Invoice Number. "
            "When the invoice is uploaded. Then the invoice should be accepted."
        )
        self.assertNotIn("Excel upload", rewritten)
        self.assertNotIn("When the invoice is uploaded", rewritten)
        self.assertIn("When the copied form is updated", rewritten)
        self.assertIn("Then Update should succeed.", rewritten)

    def test_preconditions_column_is_copy_story_not_upload(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "All Testcases"
            ws.append(
                [
                    "Test Case ID",
                    "Preconditions",
                    "Steps of Test case",
                    "Expected Results",
                    "Testcase Title",
                ]
            )
            ws.append(
                [
                    "TC-1",
                    "Covoro Excel template is mapped for upload",
                    "Upload the Excel file. When the invoice is uploaded",
                    "Then the invoice should be rejected with an error.",
                    "Excel upload · Covoro | Tax Category",
                ]
            )
            wb.save(src)
            wb.close()
            dest = Workbook()
            dest_ws = dest.active
            dest_ws.title = "Conditional"
            copy_kept_rows(str(src), {"TC-1"}, dest_ws, sheet_name="Conditional")
            self.assertEqual(dest_ws.cell(2, 2).value, COPY_INVOICE_PRECONDITIONS)
            self.assertTrue(COPY_INVOICE_PRECONDITIONS.startswith("1. "))
            self.assertIn("\n2. ", COPY_INVOICE_PRECONDITIONS)
            self.assertNotIn("Given ", COPY_INVOICE_PRECONDITIONS)
            self.assertIn("Create Copy", dest_ws.cell(2, 2).value)
            self.assertNotIn("Excel template", dest_ws.cell(2, 2).value)
            self.assertEqual(
                dest_ws.cell(2, 3).value,
                copy_invoice_steps("", "Tax Category", "Conditional"),
            )
            self.assertNotIn("Upload", dest_ws.cell(2, 3).value)
            self.assertIn("the form should show an error", dest_ws.cell(2, 4).value)
            self.assertNotIn("Excel upload", dest_ws.cell(2, 5).value)

    def test_field_sheet_gets_copy_preconditions_and_steps_when_source_lacks_them(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            _matrix_xlsx(src, [("TC-1", "Invoice Number", "An empty Invoice Number")])
            dest = Workbook()
            dest_ws = dest.active
            copy_kept_rows(str(src), {"TC-1"}, dest_ws, sheet_name="Field")
            headers = [dest_ws.cell(1, c).value for c in range(1, 8)]
            self.assertIn("Preconditions", headers)
            self.assertIn("Steps of Test case", headers)
            self.assertIn("Expected Results", headers)
            pre_col = headers.index("Preconditions") + 1
            steps_col = headers.index("Steps of Test case") + 1
            self.assertEqual(dest_ws.cell(2, pre_col).value, COPY_INVOICE_PRECONDITIONS)
            steps = dest_ws.cell(2, steps_col).value
            self.assertIn("Create Copy", steps)
            self.assertIn("Invoice Number", steps)
            self.assertNotIn("Upload", steps)

    def test_all_three_sheets_use_copy_preconditions_and_steps(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            field = Path(tmp) / "field.xlsx"
            formula = Path(tmp) / "formula.xlsx"
            conditional = Path(tmp) / "conditional.xlsx"
            _matrix_xlsx(field, [("F-1", "Invoice Number", "f")])
            _matrix_xlsx(formula, [("P-1", "Item Net Price", "p")])
            _matrix_xlsx(conditional, [("C-1", "Tax Category", "c")])
            out = Path(tmp) / "out.xlsx"
            write_combined_workbook(
                {
                    "outputPath": str(out),
                    "sheets": [
                        {"name": "Field", "sourcePath": str(field), "keepIds": ["F-1"]},
                        {"name": "Formula", "sourcePath": str(formula), "keepIds": ["P-1"]},
                        {"name": "Conditional", "sourcePath": str(conditional), "keepIds": ["C-1"]},
                    ],
                }
            )
            wb = load_workbook(out, read_only=True, data_only=True)
            try:
                for name, field in (
                    ("Field", "Invoice Number"),
                    ("Formula", "Item Net Price"),
                    ("Conditional", "Tax Category"),
                ):
                    ws = wb[name]
                    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                    self.assertIn("Preconditions", headers, name)
                    self.assertIn("Steps of Test case", headers, name)
                    pre = ws.cell(2, headers.index("Preconditions") + 1).value
                    steps = ws.cell(2, headers.index("Steps of Test case") + 1).value
                    self.assertEqual(pre, COPY_INVOICE_PRECONDITIONS)
                    self.assertIn("Create Copy", steps)
                    self.assertIn(field, steps)
                    self.assertNotIn("Upload", steps)
            finally:
                wb.close()


class CopyKeptRowsTests(unittest.TestCase):
    def test_copies_keep_ids_preserves_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            _matrix_xlsx(
                src,
                [
                    ("TC-1", "Invoice Number", "first"),
                    ("TC-2", "Seller Name", "second"),
                    ("TC-3", "Tax Category", "third"),
                ],
            )
            dest = Workbook()
            ws = dest.active
            ws.title = "Field"
            result = copy_kept_rows(str(src), {"TC-1", "TC-2", "TC-3"}, ws)
            self.assertEqual(result["copied"], 3)
            self.assertEqual(result["sourceSheet"], "All Testcases")
            self.assertEqual(ws.cell(1, 1).value, "Test Case ID")
            self.assertEqual(ws.cell(2, 1).value, "TC-1")
            self.assertEqual(ws.cell(4, 1).value, "TC-3")

    def test_zero_keep_ids_writes_header_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            _matrix_xlsx(src, [("TC-1", "Invoice Number", "x")])
            dest = Workbook()
            result = copy_kept_rows(str(src), set(), dest.active)
            self.assertEqual(result["copied"], 0)
            self.assertEqual(dest.active.cell(1, 1).value, "Test Case ID")
            self.assertIsNone(dest.active.cell(2, 1).value)

    def test_missing_test_case_id_raises(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "All Testcases"
            ws.append(["Filed name", "Testcase Title"])
            ws.append(["Invoice Number", "x"])
            wb.save(src)
            wb.close()
            dest = Workbook()
            with self.assertRaisesRegex(ValueError, "Test Case ID"):
                copy_kept_rows(str(src), {"TC-1"}, dest.active)

    def test_write_combined_three_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            field = Path(tmp) / "field.xlsx"
            formula = Path(tmp) / "formula.xlsx"
            conditional = Path(tmp) / "conditional.xlsx"
            _matrix_xlsx(field, [("F-1", "Invoice Number", "f")])
            _matrix_xlsx(formula, [("P-1", "Item Net Price", "p")])
            _matrix_xlsx(conditional, [("C-1", "Tax Category", "c")])
            out = Path(tmp) / "out.xlsx"
            plan = {
                "outputPath": str(out),
                "sheets": [
                    {"name": "Field", "sourcePath": str(field), "keepIds": ["F-1"]},
                    {"name": "Formula", "sourcePath": str(formula), "keepIds": ["P-1"]},
                    {"name": "Conditional", "sourcePath": str(conditional), "keepIds": ["C-1"]},
                ],
            }
            write_combined_workbook(plan)
            wb = load_workbook(out, read_only=True, data_only=True)
            try:
                self.assertEqual(wb.sheetnames, ["Field", "Formula", "Conditional"])
                self.assertEqual(wb["Field"].cell(2, 1).value, "F-1")
                self.assertEqual(wb["Formula"].cell(2, 1).value, "P-1")
                self.assertEqual(wb["Conditional"].cell(2, 1).value, "C-1")
            finally:
                wb.close()


if __name__ == "__main__":
    raise SystemExit(unittest.main())
