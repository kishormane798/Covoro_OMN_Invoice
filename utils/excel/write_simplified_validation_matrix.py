#!/usr/bin/env python3
"""Copy kept FullMatrix rows into one Simplified catalog workbook (3 sheets)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

def _emit_error(payload: dict) -> None:
    line = json.dumps(payload)
    print(line)
    print(line, file=sys.stderr)


try:
    import openpyxl
    from openpyxl import Workbook
except ImportError as exc:  # pragma: no cover
    _emit_error({"ok": False, "error": f"openpyxl required: {exc}"})
    sys.exit(1)


def _source_sheet(wb: openpyxl.Workbook):
    if "All Testcases" in wb.sheetnames:
        return wb["All Testcases"]
    return wb[wb.sheetnames[0]]


def _header_index(headers: list[str], name: str) -> int | None:
    want = name.strip().lower()
    for i, h in enumerate(headers):
        if (h or "").strip().lower() == want:
            return i
    return None


def copy_kept_rows(source_path: str, keep_ids: set[str], dest_ws) -> dict:
    path = Path(source_path)
    if not path.is_file():
        raise FileNotFoundError(f"Source matrix not found: {path}")

    src_wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        src_ws = _source_sheet(src_wb)
        rows_iter = src_ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise ValueError(f"Empty sheet in {path}") from exc
        headers = ["" if c is None else str(c) for c in header_row]
        id_idx = _header_index(headers, "Test Case ID")
        if id_idx is None:
            raise ValueError(f"Test Case ID column missing on sheet {src_ws.title}")

        for col, header in enumerate(headers, start=1):
            dest_ws.cell(1, col, header)

        copied = 0
        dest_row = 2
        for row in rows_iter:
            if not row or id_idx >= len(row):
                continue
            tc_id = "" if row[id_idx] is None else str(row[id_idx]).strip()
            if not tc_id or tc_id not in keep_ids:
                continue
            for col, value in enumerate(row, start=1):
                dest_ws.cell(dest_row, col, value)
            copied += 1
            dest_row += 1
        return {"copied": copied, "sourceSheet": src_ws.title}
    finally:
        src_wb.close()


def write_combined_workbook(plan: dict) -> dict:
    output_path = Path(plan["outputPath"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    dest_wb = Workbook()
    default = dest_wb.active
    dest_wb.remove(default)
    sheets_out = []
    for sheet in plan["sheets"]:
        ws = dest_wb.create_sheet(sheet["name"])
        keep_ids = {str(x).strip() for x in sheet.get("keepIds") or [] if str(x).strip()}
        result = copy_kept_rows(sheet["sourcePath"], keep_ids, ws)
        sheets_out.append(
            {
                "name": sheet["name"],
                "copied": result["copied"],
                "sourceSheet": result["sourceSheet"],
            }
        )
    dest_wb.save(output_path)
    dest_wb.close()
    return {"ok": True, "outputPath": str(output_path), "sheets": sheets_out}


def main() -> int:
    if len(sys.argv) < 2:
        _emit_error({"ok": False, "error": "Usage: write_simplified_validation_matrix.py <plan.json>"})
        return 1
    plan_path = Path(sys.argv[1])
    if not plan_path.is_file():
        _emit_error({"ok": False, "error": f"Plan not found: {plan_path}"})
        return 1
    try:
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        payload = write_combined_workbook(plan)
    except Exception as exc:
        _emit_error({"ok": False, "error": str(exc)})
        return 1
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
