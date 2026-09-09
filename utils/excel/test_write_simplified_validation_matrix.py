from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

_UTILS = Path(__file__).resolve().parent
if str(_UTILS) not in sys.path:
    sys.path.insert(0, str(_UTILS))

from write_simplified_validation_matrix import (  # noqa: E402
    copy_kept_rows,
    write_combined_workbook,
)


def _matrix_xlsx(path: Path, rows: list[tuple[str, str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Testcases"
    ws.append(["Test Case ID", "Filed name", "Testcase Title"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    wb.close()


class CopyKeptRowsTests(unittest.TestCase):
    def test_copies_only_keep_ids_and_preserves_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            _matrix_xlsx(
                src,
                [
                    ("TC-1", "Invoice Number", "keep first"),
                    ("TC-2", "Seller Address Line 1", "drop"),
                    ("TC-3", "Tax Category", "keep second"),
                ],
            )
            dest = Workbook()
            ws = dest.active
            ws.title = "Field"
            result = copy_kept_rows(str(src), {"TC-1", "TC-3"}, ws)
            self.assertEqual(result["copied"], 2)
            self.assertEqual(result["sourceSheet"], "All Testcases")
            self.assertEqual(ws.cell(1, 1).value, "Test Case ID")
            self.assertEqual(ws.cell(2, 1).value, "TC-1")
            self.assertEqual(ws.cell(3, 1).value, "TC-3")
            self.assertIsNone(ws.cell(4, 1).value)

    def test_zero_keep_ids_writes_header_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            _matrix_xlsx(src, [("TC-1", "Invoice Number", "x")])
            dest = Workbook()
            ws = dest.active
            result = copy_kept_rows(str(src), set(), ws)
            self.assertEqual(result["copied"], 0)
            self.assertEqual(ws.cell(1, 1).value, "Test Case ID")
            self.assertIsNone(ws.cell(2, 1).value)

    def test_missing_test_case_id_raises(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "All Testcases"
            ws.append(["Filed name", "Testcase Title"])
            ws.append(["Invoice Number", "x"])
            wb.save(src)
            wb.close()
            dest = Workbook()
            with self.assertRaisesRegex(ValueError, "Test Case ID"):
                copy_kept_rows(str(src), {"TC-1"}, dest.active)

    def test_write_combined_three_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            field = Path(tmp) / "field.xlsx"
            formula = Path(tmp) / "formula.xlsx"
            conditional = Path(tmp) / "conditional.xlsx"
            _matrix_xlsx(field, [("F-1", "Invoice Number", "f")])
            _matrix_xlsx(formula, [("P-1", "Item Net Price", "p")])
            _matrix_xlsx(conditional, [("C-1", "Tax Category", "c")])
            out = Path(tmp) / "out.xlsx"
            plan = {
                "outputPath": str(out),
                "sheets": [
                    {"name": "Field", "sourcePath": str(field), "keepIds": ["F-1"]},
                    {"name": "Formula", "sourcePath": str(formula), "keepIds": ["P-1"]},
                    {"name": "Conditional", "sourcePath": str(conditional), "keepIds": ["C-1"]},
                ],
            }
            write_combined_workbook(plan)
            wb = load_workbook(out, read_only=True, data_only=True)
            try:
                self.assertEqual(wb.sheetnames, ["Field", "Formula", "Conditional"])
                self.assertEqual(wb["Field"].cell(2, 1).value, "F-1")
                self.assertEqual(wb["Formula"].cell(2, 1).value, "P-1")
                self.assertEqual(wb["Conditional"].cell(2, 1).value, "C-1")
            finally:
                wb.close()


if __name__ == "__main__":
    raise SystemExit(unittest.main())
