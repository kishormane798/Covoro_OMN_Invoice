#!/usr/bin/env python3
"""Copy kept FullMatrix rows into one Copy-invoice catalog workbook (3 sheets).

Every sheet (Field, Formula, Conditional) gets Copy Invoice Preconditions / Steps /
Expected. Source Excel-upload narrative is not kept. Test Case IDs and field
columns stay the same.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl import Workbook
except ImportError as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"openpyxl required: {exc}"}))
    sys.exit(1)

COPY_INVOICE_PRECONDITIONS = (
    "1. Log into the COVORO OMAN Portal.\n"
    "2. Open Business Dashboard for the E Invoice product for a particular Business.\n"
    "3. Have an invoice in a Copy-eligible status "
    "(Ready to Submit, Delivered, Delivered to C3, Delivered to C5, Disregard, or Submit).\n"
    "4. The invoice is not in Error in records or Submission Error.\n"
    "5. Click Create Copy from the action dropdown.\n"
    "6. Confirm the popup: Are you sure you want to create copy of this Invoice?\n"
    "7. Click Yes.\n"
    "8. The Create E Invoice form opens with Invoice Number and Invoice Issue Date blank.\n"
    "9. The rest of the details match the original invoice."
)

REQUIRED_COPY_COLUMNS = ("Preconditions", "Steps of Test case", "Expected Results")

_TITLE_HEADERS = {
    "testcase title",
    "test case title",
    "test cases title",
}

_FIELD_HEADERS = {"filed name", "field name"}

_DESC_HEADERS = {"test cases description", "description"}

_UPLOAD_TO_COPY = (
    ("Excel upload · Covoro | ", ""),
    ("Excel upload · Covoro", "Copy Invoice"),
    ("Excel upload", "Copy Invoice"),
    ("excel upload", "Copy Invoice"),
    ("When the invoice is uploaded", "When the copied form is updated"),
    ("when the invoice is uploaded", "when the copied form is updated"),
    ("Then the invoice should be accepted.", "Then Update should succeed."),
    ("Then the invoice should be rejected with an error.", "Then the form should show an error."),
    ("the invoice should be accepted", "Update should succeed"),
    ("the invoice should be rejected with an error", "the form should show an error"),
    ("Upload the Excel file", "Update the copied invoice form"),
    ("upload the Excel file", "update the copied invoice form"),
    ("Upload the Excel", "Update the copied invoice form"),
    ("upload the Excel", "update the copied invoice form"),
    ("Upload the invoice", "Update the copied form"),
    ("upload the invoice", "update the copied form"),
    ("Covoro Excel template", "copied invoice form"),
    ("Covoro template", "copied invoice form"),
    ("Excel template", "copied invoice form"),
    ("error file", "form error"),
    ("Error file", "form error"),
)


def rewrite_copy_invoice_narrative(text: str) -> str:
    if text is None:
        return ""
    rewritten = str(text)
    for old, new in _UPLOAD_TO_COPY:
        rewritten = rewritten.replace(old, new)
    return rewritten.strip()


def copy_invoice_steps(field: str, title: str, sheet_name: str) -> str:
    display_title = rewrite_copy_invoice_narrative(title)
    field_part = (field or "").strip() or display_title or "the field under test"
    kind = (sheet_name or "Field").strip()
    detail = f": {display_title}" if display_title and display_title != field_part else ""
    return (
        "1. Click Create Copy from the invoice action dropdown.\n"
        "2. Confirm Yes on Are you sure you want to create copy of this Invoice?\n"
        "3. Confirm Invoice Number and Invoice Issue Date are blank, and the rest of the details match the original invoice.\n"
        "4. Enter a new Invoice Number and Invoice Issue Date.\n"
        f"5. Apply this {kind} case on the copied form for {field_part}{detail}.\n"
        "6. Click Update."
    )


def copy_invoice_expected(original: str, title: str, polarity: str = "") -> str:
    blob = f"{polarity} {original} {title}".lower()
    if any(
        hint in blob
        for hint in (
            "reject",
            "error file",
            "form should show an error",
            "negative",
            "should be rejected",
        )
    ):
        return "Then the form should show an error."
    if any(hint in blob for hint in ("accepted", "succeed", "positive")):
        return "Then Update should succeed."
    rewritten = rewrite_copy_invoice_narrative(original)
    if "form should show an error" in rewritten.lower():
        return "Then the form should show an error."
    if "update should succeed" in rewritten.lower():
        return "Then Update should succeed."
    return rewritten or "Then Update should succeed."


def _norm_header(header: str) -> str:
    return " ".join((header or "").strip().lower().split())


def _ensure_copy_headers(headers: list[str]) -> list[str]:
    existing = {_norm_header(h) for h in headers}
    out = list(headers)
    for col in REQUIRED_COPY_COLUMNS:
        if _norm_header(col) not in existing:
            out.append(col)
    return out


def _cell_by_header(values: dict[str, object], names: set[str]) -> str:
    for header, value in values.items():
        if _norm_header(header) in names:
            return "" if value is None else str(value).strip()
    return ""


def _source_sheet(wb: openpyxl.Workbook):
    if "All Testcases" in wb.sheetnames:
        return wb["All Testcases"]
    return wb[wb.sheetnames[0]]


def _header_index(headers: list[str], name: str) -> int | None:
    want = name.strip().lower()
    for i, h in enumerate(headers):
        if (h or "").strip().lower() == want:
            return i
    return None


def _value_for_header(
    header: str,
    values: dict[str, object],
    sheet_name: str,
) -> object:
    name = _norm_header(header)
    field = _cell_by_header(values, _FIELD_HEADERS)
    title = _cell_by_header(values, _TITLE_HEADERS)
    polarity = _cell_by_header(values, {"polarity"})
    original_expected = _cell_by_header(values, {"expected results", "expected result"})
    if name == "preconditions":
        return COPY_INVOICE_PRECONDITIONS
    if name in {"steps of test case", "steps"}:
        return copy_invoice_steps(field, title, sheet_name)
    if name in {"expected results", "expected result"}:
        return copy_invoice_expected(original_expected, title, polarity)
    if name in _TITLE_HEADERS or name in _DESC_HEADERS:
        raw = values.get(header)
        return rewrite_copy_invoice_narrative("" if raw is None else str(raw))
    return values.get(header)


def copy_kept_rows(
    source_path: str,
    keep_ids: set[str],
    dest_ws,
    sheet_name: str = "Field",
) -> dict:
    path = Path(source_path)
    if not path.is_file():
        raise FileNotFoundError(f"Source matrix not found: {path}")

    src_wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        src_ws = _source_sheet(src_wb)
        rows_iter = src_ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise ValueError(f"Empty sheet in {path}") from exc
        source_headers = ["" if c is None else str(c) for c in header_row]
        id_idx = _header_index(source_headers, "Test Case ID")
        if id_idx is None:
            raise ValueError(f"Test Case ID column missing on sheet {src_ws.title}")

        dest_headers = _ensure_copy_headers(source_headers)
        for col, header in enumerate(dest_headers, start=1):
            dest_ws.cell(1, col, header)

        copied = 0
        dest_row = 2
        for row in rows_iter:
            if not row or id_idx >= len(row):
                continue
            tc_id = "" if row[id_idx] is None else str(row[id_idx]).strip()
            if not tc_id or tc_id not in keep_ids:
                continue
            values = {
                source_headers[i]: (row[i] if i < len(row) else None)
                for i in range(len(source_headers))
            }
            for col, header in enumerate(dest_headers, start=1):
                dest_ws.cell(dest_row, col, _value_for_header(header, values, sheet_name))
            copied += 1
            dest_row += 1
        return {"copied": copied, "sourceSheet": src_ws.title}
    finally:
        src_wb.close()


def write_combined_workbook(plan: dict) -> dict:
    output_path = Path(plan["outputPath"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    dest_wb = Workbook()
    default = dest_wb.active
    dest_wb.remove(default)
    sheets_out = []
    for sheet in plan["sheets"]:
        ws = dest_wb.create_sheet(sheet["name"])
        keep_ids = {str(x).strip() for x in sheet.get("keepIds") or [] if str(x).strip()}
        result = copy_kept_rows(
            sheet["sourcePath"],
            keep_ids,
            ws,
            sheet_name=str(sheet.get("name") or "Field"),
        )
        sheets_out.append(
            {
                "name": sheet["name"],
                "copied": result["copied"],
                "sourceSheet": result["sourceSheet"],
            }
        )
    dest_wb.save(output_path)
    dest_wb.close()
    return {"ok": True, "outputPath": str(output_path), "sheets": sheets_out}


def main() -> int:
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "Usage: write_copy_validation_matrix.py <plan.json>"}))
        return 1
    plan_path = Path(sys.argv[1])
    if not plan_path.is_file():
        print(json.dumps({"ok": False, "error": f"Plan not found: {plan_path}"}))
        return 1
    try:
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        payload = write_combined_workbook(plan)
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        return 1
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
